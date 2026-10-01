"""Sample policy and report documents (fictional, for the Sahyadripur municipal corporation). Used by RAG in Phase 7."""

from datetime import datetime
from pathlib import Path

from docx import Document
from fpdf import FPDF

from scripts.datagen.constants import CITY, DEPARTMENTS

Section = tuple[str, list[str]]

SLA_LINES = [f"{d.name}: {d.sla_hours} hours for complaints in its category." for d in DEPARTMENTS]

DOCS: list[dict[str, object]] = [
    {
        "file": "citizen_grievance_redressal_policy_2024.pdf", "title": "Citizen Grievance Redressal Policy 2024",
        "doc_type": "policy",
        "sections": [
            ("1. Purpose", [f"This policy sets out how the {CITY} Municipal Corporation receives, routes, resolves and reviews citizen complaints. It applies to all departments and all 20 wards. The corporation commits to a transparent, time-bound grievance process in which every complaint receives a tracking number and every citizen can follow its status."]),
            ("2. Service level agreements", ["Each complaint category has a maximum resolution time counted from the moment the complaint is registered:"] + SLA_LINES + ["A complaint that is not resolved within its SLA is recorded as an SLA breach and is escalated automatically."]),
            ("3. Escalation matrix", ["Level 1: Ward Officer, at registration. Level 2: Department Head, when 100 percent of the SLA has elapsed. Level 3: Additional Commissioner, when 200 percent of the SLA has elapsed. Level 4: Municipal Commissioner, when 300 percent of the SLA has elapsed or when the complaint concerns a public health risk. Complaints reported by more than 25 citizens within 7 days about the same location are treated as a cluster and escalated directly to Level 2."]),
            ("4. Feedback and reopening", ["After resolution the citizen is invited to rate the service from 1 to 5. A rating of 2 or lower reopens the complaint for review by the Ward Officer within 48 hours. Departments with an average rating below 3.0 in a quarter must submit an improvement plan to the Commissioner."]),
            ("5. Data use and privacy", ["Citizen names and phone numbers are used only to process the complaint. They are masked in analytical reports and dashboards and must never be shared with third-party services. Analytical summaries use anonymised text only."]),
            ("6. Review", ["The Commissioner reviews grievance performance monthly using resolution rate, average resolution time, SLA compliance, repeat complaint rate and citizen satisfaction. Wards in the bottom quartile for SLA compliance for two consecutive months receive a corrective action review."]),
        ],
    },
    {
        "file": "solid_waste_management_policy.pdf", "title": "Solid Waste Management Policy",
        "doc_type": "policy",
        "sections": [
            ("1. Objectives", ["Achieve door-to-door collection in every ward, reach 95 percent source segregation by 2027, and keep uncollected garbage complaints below 4 per 10,000 residents per month."]),
            ("2. Collection standards", ["Residential areas are served daily between 6 am and 11 am. Markets and commercial streets are served twice daily. A missed collection must be reported by the ward supervisor and cleared within 24 hours, which is the SLA for garbage complaints."]),
            ("3. Festival and seasonal surge plan", ["Garbage volumes rise 20 to 30 percent during the Ganesh festival, Navratri and Diwali (September to November). Departments must add temporary vehicles and shifts from the first week of September, and pre-position additional bins in Old City, Market Yard and Gandhi Chowk."]),
            ("4. Contractor performance", ["Collection contractors are evaluated monthly on missed pickups, response time to complaints and citizen ratings. Two consecutive months of missed-pickup rates above 3 percent trigger a penalty of 5 percent of the monthly payment. Work stoppages must be notified to the Commissioner within 24 hours and a fall-back crew deployed within 48 hours."]),
            ("5. Penalties and enforcement", ["Open dumping and burning of waste are punishable by a fine. Ward officers carry out weekly inspections of vacant plots that have previously generated dumping complaints."]),
        ],
    },
    {
        "file": "monsoon_preparedness_and_drainage_plan_2025.pdf", "title": "Monsoon Preparedness and Drainage Plan 2025",
        "doc_type": "policy",
        "sections": [
            ("1. Background", ["Drainage and waterlogging complaints rise three to four times between June and September. Historic hotspots include Market Yard, Railway Colony, Sahyadri Heights and Riverside. Response performance in monsoon months has historically been 30 percent slower than in the dry season."]),
            ("2. Pre-monsoon works", ["Desilting of all major storm drains must be completed by 31 May. The Drainage and Sewerage Department will clean 100 percent of chronic waterlogging points and repair all broken manhole covers reported since January. A ward-wise readiness checklist is signed by each Ward Officer."]),
            ("3. Response during the monsoon", ["A 24-hour control room operates from 1 June to 30 September. Target response time for flooding complaints is 6 hours. Mobile pumping crews are stationed in wards W07, W09 and W10. When rainfall exceeds 65 mm in 24 hours, all leave in the Drainage and Roads departments is suspended."]),
            ("4. Resources", ["Additional crews equivalent to 30 crew-days per 100 forecast drainage complaints are to be mobilised. The forecasting unit issues a monthly complaint forecast so that crews and pumps can be planned one to three months ahead."]),
            ("5. Post-monsoon review", ["Within 30 days of the monsoon ending, each ward reports the number of flooding events, hours of disruption and complaints resolved outside SLA. Lessons are used to update the plan for the following year."]),
        ],
    },
    {
        "file": "road_maintenance_and_pothole_repair_sop.pdf", "title": "Road Maintenance and Pothole Repair SOP",
        "doc_type": "policy",
        "sections": [
            ("1. Scope", ["This procedure covers repair of potholes, damaged carriageways, footpaths and speed breakers on roads maintained by the corporation, including the Ring Road."]),
            ("2. Repair timelines", ["Potholes reported on arterial roads and the Ring Road must be repaired within 5 days (the roads SLA is 120 hours). Potholes causing an accident or near a school or hospital are repaired within 48 hours. In the monsoon, temporary cold-mix patching is used and permanent repair follows within 15 days of the rains stopping."]),
            ("3. Quality and resurfacing", ["Every road is scheduled for resurfacing on a five-year cycle. Roads with three or more repeat pothole complaints within 90 days must be moved up in the resurfacing schedule. Contractors provide a 3-year defect liability period. Utility companies must restore any road surface they dig up within 7 days."]),
            ("4. Monitoring", ["Ward engineers inspect at least 20 percent of completed repairs each month. Complaint clusters along the same road segment are reviewed monthly by the Roads and Maintenance Department. A sudden rise in pothole complaints in any ward is treated as an early warning for resurfacing demand."]),
        ],
    },
    {
        "file": "water_supply_service_charter.pdf", "title": "Water Supply Service Charter",
        "doc_type": "policy",
        "sections": [
            ("1. Service commitments", ["The Water Supply Department commits to supply of at least 2 hours twice daily in every ward, minimum residual pressure at the tap, and resolution of supply complaints within 48 hours."]),
            ("2. Water quality", ["Samples are tested at least weekly at source points and monthly at consumer taps in every ward. Residual chlorine must be maintained above 0.2 mg per litre. Results are published on the corporation website."]),
            ("3. Contamination response protocol", ["When three or more complaints about discoloured, foul-smelling or contaminated water arrive from the same locality within 48 hours, the Ward Officer must alert the Public Health Department, issue a boil-water advisory within 12 hours, arrange alternate supply by tanker and complete a pipeline inspection within 24 hours."]),
            ("4. Summer demand", ["Water complaints increase 40 to 70 percent in April, May and June. Tanker capacity must be increased in outer wards, where piped coverage is lower than in central wards."]),
        ],
    },
    {
        "file": "public_health_vector_control_guidelines.pdf", "title": "Public Health and Vector Control Guidelines",
        "doc_type": "policy",
        "sections": [
            ("1. Purpose", ["To reduce the incidence of vector-borne disease, especially dengue and malaria, through routine and outbreak control measures."]),
            ("2. Routine measures", ["Larval surveys are done fortnightly in every ward from June to November. Fogging cycles run every 15 days in wards with high mosquito density and weekly in wards reporting fever clusters."]),
            ("3. Outbreak triggers", ["An outbreak alert is declared when a ward records more than 30 percent growth in public health complaints week on week for three consecutive weeks, or when five or more dengue cases are reported from one locality. The alert mobilises fogging teams within 24 hours, door-to-door surveys and notification to the district health office."]),
            ("4. Seasonality", ["Public health complaints peak between July and October. Additional teams should be deployed from the first week of July in Ambedkar Nagar, Tilak Road and other wards with historically high counts."]),
        ],
    },
    {
        "file": "annual_municipal_performance_report_2024_25.docx", "title": "Annual Municipal Performance Report 2024-25",
        "doc_type": "report",
        "sections": [
            ("Executive summary", [f"The {CITY} Municipal Corporation registered a rising volume of citizen complaints in 2024-25, about 8 percent more than the previous year. Overall SLA compliance was 61 percent. Garbage, street lighting, electricity and water complaints were resolved fastest, while drainage, roads and encroachment remained the weakest departments against their SLAs."]),
            ("Key achievements", ["Door-to-door garbage collection reached all wards. LED conversion of street lights was completed in 8 wards. The mobile app became the largest complaint channel at about 38 percent of all complaints. The grievance redressal policy was adopted, introducing a four-level escalation matrix."]),
            ("Challenges", ["Ambedkar Nagar (W05), Ganesh Peth (W12) and Hanuman Wadi (W17) had resolution times nearly double those of the best wards. Drainage complaints surged in July and August, and the June to September period saw the largest backlogs. Staff shortages in the Roads department and rising pothole complaints on the Ring Road are a growing concern. A garbage collection breakdown in March 2025 affected Ganesh Peth and Kasturba Nagar for about two weeks."]),
            ("Recommendations", ["Add crews in the underperforming wards before the next monsoon, bring forward desilting to April, strengthen contractor penalties in solid waste management, and use complaint forecasts to plan seasonal resources. The Commissioner should review wards in the bottom quartile monthly."]),
        ],
        "table": [["Department", "Complaint SLA (hours)", "SLA compliance"], ["Solid Waste Management", "24", "84%"], ["Street Lighting", "72", "76%"], ["Electricity", "24", "75%"], ["Water Supply", "48", "74%"], ["Public Health", "48", "57%"], ["Drainage & Sewerage", "72", "43%"], ["Encroachment Removal", "168", "38%"], ["Roads & Maintenance", "120", "36%"]],
    },
    {
        "file": "street_lighting_and_energy_efficiency_policy.docx", "title": "Street Lighting and Energy Efficiency Policy",
        "doc_type": "policy",
        "sections": [
            ("1. Goals", ["Convert all street lights to LED by 2027, reduce lighting energy use by 40 percent, and keep the share of non-working lights below 2 percent at any time."]),
            ("2. Maintenance standards", ["Street light complaints must be resolved within 72 hours. Ward electricians inspect every lane at least once a month. A lane with more than three non-working lights is treated as a priority."]),
            ("3. Outage cluster protocol", ["When ten or more street light complaints are received from the same lane or neighbourhood within 7 days, the Street Lighting Department must inspect the feeder and distribution panel within 24 hours, since the cause is usually a cable or panel fault rather than individual lamps. Police are informed of dark stretches in areas that report safety concerns."]),
            ("4. Energy efficiency", ["Central controllers with timers and dimming are installed on arterial roads. Daytime-on lights are reported and fixed within 48 hours to avoid energy waste."]),
        ],
    },
]


def _write_pdf(path: Path, doc: dict[str, object]) -> None:
    pdf = FPDF()
    pdf.set_creation_date(datetime(2025, 4, 1))
    pdf.set_title(str(doc["title"]))
    pdf.set_margins(18, 18, 18)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 8, str(doc["title"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "I", 10)
    pdf.multi_cell(0, 6, f"{CITY} Municipal Corporation", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    for heading, paras in doc["sections"]:  # type: ignore[attr-defined]
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(0, 7, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for p in paras:
            pdf.multi_cell(0, 6, p, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        pdf.ln(2)
    pdf.output(str(path))


def _write_docx(path: Path, doc: dict[str, object]) -> None:
    d = Document()
    d.core_properties.created = datetime(2025, 4, 1)
    d.core_properties.title = str(doc["title"])
    d.add_heading(str(doc["title"]), 0)
    d.add_paragraph(f"{CITY} Municipal Corporation")
    for heading, paras in doc["sections"]:  # type: ignore[attr-defined]
        d.add_heading(heading, 1)
        for p in paras:
            d.add_paragraph(p)
    table_rows = doc.get("table")
    if table_rows:
        d.add_heading("Department SLA compliance", 1)
        t = d.add_table(rows=len(table_rows), cols=len(table_rows[0]))  # type: ignore[arg-type,index]
        t.style = "Table Grid"
        for i, row in enumerate(table_rows):  # type: ignore[arg-type]
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    d.save(str(path))


def write_policy_documents(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for doc in DOCS:
        path = out_dir / str(doc["file"])
        (_write_docx if path.suffix == ".docx" else _write_pdf)(path, doc)
        paths.append(path)
    return paths
