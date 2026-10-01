"""Policy/report documents: text extraction (PDF, DOCX, TXT/MD), chunking for RAG, and storage of the chunks.

Chunk embeddings are left NULL here; they are computed by the embedding job (Phase 3 / RAG in Phase 7).
"""

import hashlib
import io
import re
from dataclasses import dataclass
from typing import Any

from docx import Document as DocxDocument
from pypdf import PdfReader
from sqlalchemy import insert, select
from sqlalchemy.orm import Session

from app.db.models import Document, DocumentChunk
from app.ingestion.normalize import clean_text

ALLOWED_DOC_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
CHUNK_CHARS = 900
CHUNK_OVERLAP = 150
MAX_CHARS = 2_000_000


class DocumentError(ValueError):
    pass


@dataclass
class Chunk:
    index: int
    page: int | None
    content: str


def extract_pages(data: bytes, filename: str) -> list[tuple[int | None, str]]:
    """Return [(page number or None, text)]. Raises DocumentError for unreadable or empty files."""
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(data))
            pages = [(i + 1, (p.extract_text() or "")) for i, p in enumerate(reader.pages)]
        elif name.endswith(".docx"):
            doc = DocxDocument(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                parts += [" | ".join(c.text for c in row.cells) for row in table.rows]
            pages = [(None, "\n".join(parts))]
        elif name.endswith((".txt", ".md")):
            pages = [(None, data.decode("utf-8", errors="replace"))]
        else:
            raise DocumentError(f"Unsupported document type; allowed: {', '.join(sorted(ALLOWED_DOC_EXTENSIONS))}")
    except DocumentError:
        raise
    except Exception as exc:  # noqa: BLE001 - pypdf/python-docx raise many types for corrupt input
        raise DocumentError(f"Could not read the document: {exc}") from exc
    pages = [(p, t) for p, t in pages if t.strip()]
    if not pages:
        raise DocumentError("No extractable text found (scanned PDFs need OCR, which is not supported)")
    if sum(len(t) for _, t in pages) > MAX_CHARS:
        raise DocumentError("Document text is too large")
    return pages


def chunk_pages(pages: list[tuple[int | None, str]], size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[Chunk]:
    """Split into ~`size`-character chunks on paragraph/sentence boundaries with `overlap` characters of context."""
    chunks: list[Chunk] = []
    for page, text in pages:
        text = re.sub(r"[ \t]+", " ", text)
        paras = [clean_text(p) for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
        sentences: list[str] = []
        for p in paras:
            sentences += [s for s in re.split(r"(?<=[.!?])\s+", p) if s]
        current = ""
        for sent in sentences:
            if len(sent) > size:  # very long sentence: hard split
                pieces = [sent[i : i + size] for i in range(0, len(sent), size - overlap)]
            else:
                pieces = [sent]
            for piece in pieces:
                if current and len(current) + len(piece) + 1 > size:
                    chunks.append(Chunk(len(chunks), page, current))
                    current = current[-overlap:].split(" ", 1)[-1] if overlap else ""
                current = f"{current} {piece}".strip()
        if current:
            chunks.append(Chunk(len(chunks), page, current))
    return chunks


def content_hash(pages: list[tuple[int | None, str]]) -> str:
    return hashlib.sha256("\n".join(clean_text(t) for _, t in pages).encode("utf-8")).hexdigest()


def ingest_document(
    db: Session, data: bytes, filename: str, *, title: str, doc_type: str = "policy", user_id: int | None = None,
    storage_bucket: str | None = None, storage_path: str | None = None,
) -> tuple[Document, bool]:
    """Extract, chunk and store a document. Returns (document, created); created=False if identical text exists."""
    pages = extract_pages(data, filename)
    h = content_hash(pages)
    existing = db.scalar(select(Document).where(Document.content_hash == h))
    if existing is not None:
        return existing, False
    chunks = chunk_pages(pages)
    doc = Document(
        title=title, doc_type=doc_type, source_filename=filename, content_hash=h, num_chunks=len(chunks),
        uploaded_by=user_id, storage_bucket=storage_bucket, storage_path=storage_path,
    )
    db.add(doc)
    db.flush()
    rows: list[dict[str, Any]] = [
        {"document_id": doc.id, "chunk_index": c.index, "page": c.page, "content": c.content} for c in chunks
    ]
    for i in range(0, len(rows), 500):
        db.execute(insert(DocumentChunk.__table__).values(rows[i : i + 500]))
    db.commit()
    return doc, True
