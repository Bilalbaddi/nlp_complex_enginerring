"""Read uploaded tabular files (CSV / Excel / JSON / JSON Lines) into a list of plain dict rows."""

import csv
import io
import json
import re
from typing import Any

from openpyxl import load_workbook

MAX_ROWS = 200_000
ALLOWED_TABLE_EXTENSIONS = {".csv", ".xlsx", ".xlsm", ".json", ".jsonl", ".ndjson"}

# header synonyms -> canonical field names
HEADER_ALIASES = {
    "complaint_id": "external_id", "id": "external_id", "ticket_id": "external_id", "reference": "external_id",
    "complaint": "description", "complaint_text": "description", "details": "description", "message": "description",
    "date": "submitted_at", "created_at": "submitted_at", "timestamp": "submitted_at", "complaint_date": "submitted_at", "registered_at": "submitted_at", "datetime": "submitted_at",
    "lat": "latitude", "lon": "longitude", "lng": "longitude", "long": "longitude",
    "ward_name": "ward", "ward_no": "ward", "ward_code": "ward",
    "dept": "department", "department_name": "department",
    "type": "category", "complaint_type": "category", "issue": "category",
    "name": "citizen_name", "phone": "citizen_phone", "mobile": "citizen_phone", "contact": "citizen_phone",
    "closed_at": "resolved_at", "resolution_date": "resolved_at", "resolved_on": "resolved_at",
    "comment": "feedback_text", "comments": "feedback_text", "feedback": "feedback_text",
    "score": "rating", "stars": "rating",
    "month": "period", "date_month": "period",
    "complaint_ref": "complaint_external_id", "complaint_no": "complaint_external_id",
    "received": "complaints_received", "resolved": "complaints_resolved", "breaches": "sla_breaches", "staff": "staff_count",
}


class FileFormatError(ValueError):
    pass


def canonical_header(h: Any) -> str:
    name = re.sub(r"[^0-9a-zA-Z]+", "_", str(h or "").strip().lower()).strip("_")
    return HEADER_ALIASES.get(name, name)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise FileFormatError("Could not decode the file as UTF-8 or Windows-1252 text")


def _read_csv(data: bytes) -> list[dict[str, Any]]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel  # type: ignore[assignment]
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        header = [canonical_header(h) for h in next(reader)]
    except StopIteration:
        return []
    rows = []
    for values in reader:
        if not any(v.strip() for v in values):
            continue
        rows.append({h: v for h, v in zip(header, values) if h})
        if len(rows) > MAX_ROWS:
            raise FileFormatError(f"File has more than {MAX_ROWS:,} rows")
    return rows


def _read_xlsx(data: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        header = [canonical_header(h) for h in next(it)]
    except StopIteration:
        return []
    rows = []
    for values in it:
        if not any(v not in (None, "") for v in values):
            continue
        rows.append({h: v for h, v in zip(header, values) if h})
        if len(rows) > MAX_ROWS:
            raise FileFormatError(f"File has more than {MAX_ROWS:,} rows")
    return rows


def _read_json(data: bytes, filename: str) -> list[dict[str, Any]]:
    text = _decode(data).strip()
    try:
        if filename.lower().endswith((".jsonl", ".ndjson")):
            parsed: Any = [json.loads(line) for line in text.splitlines() if line.strip()]
        else:
            parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise FileFormatError(f"Invalid JSON: {exc.msg} (line {exc.lineno})") from exc
    if isinstance(parsed, dict):
        parsed = next((v for v in parsed.values() if isinstance(v, list)), [parsed])
    if not isinstance(parsed, list) or not all(isinstance(r, dict) for r in parsed):
        raise FileFormatError("JSON must be a list of objects")
    if len(parsed) > MAX_ROWS:
        raise FileFormatError(f"File has more than {MAX_ROWS:,} rows")
    return [{canonical_header(k): v for k, v in r.items()} for r in parsed]


def read_rows(data: bytes, filename: str) -> list[dict[str, Any]]:
    name = filename.lower()
    if name.endswith(".csv"):
        return _read_csv(data)
    if name.endswith((".xlsx", ".xlsm")):
        try:
            return _read_xlsx(data)
        except FileFormatError:
            raise
        except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for corrupt files
            raise FileFormatError(f"Could not read Excel file: {exc}") from exc
    if name.endswith((".json", ".jsonl", ".ndjson")):
        return _read_json(data, name)
    raise FileFormatError(f"Unsupported file type; allowed: {', '.join(sorted(ALLOWED_TABLE_EXTENSIONS))}")
